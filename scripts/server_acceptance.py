"""Real-file server acceptance; run inside an offline FontGuard test container."""

import json
import subprocess
import sys
import time
from pathlib import Path

import jsonschema
import pymupdf
from fontTools.ttLib import TTFont


def run(root: Path, *args: str, expected: int = 0) -> dict:
    process = subprocess.run(
        [sys.executable, "-m", "fontguard", *args],
        cwd=root,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    if process.returncode != expected:
        raise AssertionError(
            f"CLI exit {process.returncode}, expected {expected}: {process.stderr} "
            f"{process.stdout[:1000]}"
        )
    return json.loads(process.stdout) if process.stdout.lstrip().startswith("{") else {}


def main() -> None:
    import socket

    def deny_network(*args, **kwargs):
        raise AssertionError("Network attempted during offline server test")

    socket.socket.connect = deny_network
    root = Path("/acceptance")
    output = root / "results"
    output.mkdir(exist_ok=True)
    fonts = sorted((root / "fonts").glob("*.ttf"))
    assert len(fonts) == 8
    # Actual licensed upstream binaries, renamed filenames, exact SHA256 identity.
    font_report = run(root, "scan", "fonts", "--usage", "webfont", "--format", "json", "--no-cache")
    assert len(font_report["fonts"]) == 8
    assert all(f["identity"]["match_method"] == "sha256" for f in font_report["fonts"])
    (output / "actual-fonts.json").write_text(json.dumps(font_report, indent=2), encoding="utf-8")
    # Real transformed WOFF and WOFF2, whose hashes differ but names remain identifiable.
    with TTFont(fonts[0]) as font:
        for flavor in ("woff", "woff2"):
            font.flavor = flavor
            font.save(str(root / ("converted." + flavor)))
    for extension in ("woff", "woff2"):
        converted = run(
            root,
            "scan",
            f"converted.{extension}",
            "--usage",
            "webfont",
            "--format",
            "json",
            "--no-cache",
        )
        assert converted["fonts"][0]["identity"]["database_id"]
    # Create an actual PDF containing embedded font bytes and subset it.
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_font(fontname="brand", fontfile=str(fonts[0]))
        page.insert_text((40, 60), "FontGuard actual PDF 2026", fontname="brand")
        document.subset_fonts()
        document.save(str(root / "poster.pdf"))
    pdf = run(
        root,
        "scan",
        "poster.pdf",
        "--usage",
        "document_embedding",
        "--format",
        "json",
        "--no-cache",
    )
    assert pdf["fonts"] and all(f["embedded"] for f in pdf["fonts"])
    assert all(f["identity"]["database_id"] for f in pdf["fonts"])
    (output / "actual-pdf.json").write_text(json.dumps(pdf, indent=2), encoding="utf-8")
    # Create native DOCX and PPTX, then scan their actual saved ZIP/XML contents.
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches

    document = Document()
    run_text = document.add_paragraph().add_run("FontGuard server test")
    run_text.font.name = "Microsoft YaHei"
    document.save(str(root / "document.docx"))
    office = run(
        root,
        "scan",
        "document.docx",
        "--usage",
        "commercial_design",
        "--format",
        "json",
        "--no-cache",
    )
    assert any(
        f["canonical_name"] == "Microsoft YaHei" and f["risk_level"] == "REVIEW"
        for f in office["fonts"]
    )
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    run_text = text.text_frame.paragraphs[0].add_run()
    run_text.text = "Server acceptance"
    run_text.font.name = "Noto Sans SC"
    presentation.save(str(root / "slides.pptx"))
    ppt = run(
        root,
        "scan",
        "slides.pptx",
        "--usage",
        "commercial_design",
        "--format",
        "json",
        "--no-cache",
    )
    assert any(f["canonical_name"] == "Noto Sans SC" for f in ppt["fonts"])
    (output / "actual-office.json").write_text(
        json.dumps({"docx": office, "pptx": ppt}, indent=2), encoding="utf-8"
    )
    # HTML/SVG/local CSS references, policy, caching and baseline regression.
    web = root / "web"
    web.mkdir(exist_ok=True)
    reference = Path("fonts") / fonts[0].name
    # Scan root must include sibling fonts; the local alias resolves to actual bytes.
    (root / "site.css").write_text(
        f'@font-face {{font-family: "Brand"; src: url("{reference.as_posix()}")}}'
        '\nbody {font-family: "Brand", "Noto Sans SC"}',
        encoding="utf-8",
    )
    css = run(root, "scan", "site.css", "--usage", "webfont", "--format", "json", "--no-cache")
    assert any(f["detection_method"] == "css_font_source" for f in css["fonts"])
    (web / "index.html").write_text(
        '<span style="font-family: Arial">Test</span>', encoding="utf-8"
    )
    (web / "poster.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg">'
        '<text font-family="Noto Sans TC">Test</text></svg>',
        encoding="utf-8",
    )
    cold = run(root, "scan", "web", "--usage", "webfont", "--format", "json")
    warm = run(root, "scan", "web", "--usage", "webfont", "--format", "json")
    assert cold["files_scanned"] == 2 and warm["cache_hits"] == 2
    run(root, "scan", "web", "--fail-on", "review", "--format", "json", expected=1)
    run(root, "baseline", "create", "web", "--output", "baseline.json")
    run(
        root,
        "scan",
        "web",
        "--baseline",
        "baseline.json",
        "--fail-on",
        "review",
        "--format",
        "json",
    )
    (web / "new.css").write_text('body {font-family: "Server Unknown Brand"}', encoding="utf-8")
    run(
        root,
        "scan",
        "web",
        "--baseline",
        "baseline.json",
        "--fail-on",
        "unknown",
        "--format",
        "json",
        expected=1,
    )
    (web / ".fontguard.yml").write_text(
        "version: 1\npolicy:\n  denied_fonts: [Arial]\n", encoding="utf-8"
    )
    run(root, "scan", "web", "--fail-on", "high", "--format", "json", expected=1)
    (root / "broken.ttf").write_bytes(b"not a font")
    run(root, "scan", "broken.ttf", "--format", "json", expected=2)
    # Complete SARIF validation and final user-visible reports.
    run(root, "scan", "web", "--format", "sarif", "--output", "results/server.sarif", expected=1)
    schema = json.loads(Path("/opt/fontguard/tests/schemas/sarif-2.1.0.json").read_text())
    jsonschema.Draft7Validator(schema).validate(json.loads((output / "server.sarif").read_text()))
    run(root, "scan", "web", "--format", "html", "--output", "results/server.html", expected=1)
    result = {
        "status": "passed",
        "actual_fonts": len(fonts),
        "exact_hash_matches": 8,
        "formats": ["TTF", "WOFF", "WOFF2", "PDF", "DOCX", "PPTX", "CSS", "HTML", "SVG"],
        "cache": "passed",
        "baseline": "passed",
        "policy": "passed",
        "sarif_schema": "passed",
        "exit_codes": [0, 1, 2],
        "network": "disabled",
        "python": sys.version,
        "timestamp": time.time(),
    }
    (output / "acceptance.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
